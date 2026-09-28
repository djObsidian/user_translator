"""Rendering the block model into a .docx."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

from .markdown import Block, html_table_rows

# **bold**, __bold__, *italic*, _italic_, `code`
_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|(?<!\w)\*[^*\s][^*]*?\*|(?<!\w)_[^_\s][^_]*?_(?!\w)|`[^`]+`)")
_HTML_TAG = re.compile(r"<[^>]+>")


def _add_inline(par, text: str, bold: bool = False) -> None:
    text = _HTML_TAG.sub("", text)
    for tok in _INLINE.split(text):
        if not tok:
            continue
        run_bold, italic, mono = bold, False, False
        if tok.startswith(("**", "__")) and len(tok) > 4:
            tok, run_bold = tok[2:-2], True
        elif tok.startswith("`"):
            tok, mono = tok[1:-1], True
        elif tok[0] in "*_" and tok[-1] == tok[0] and len(tok) > 2:
            tok, italic = tok[1:-1], True
        run = par.add_run(tok)
        run.bold = run_bold or None
        run.italic = italic or None
        if mono:
            run.font.name = "Consolas"


def _add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            par = table.cell(i, j).paragraphs[0]
            _add_inline(par, cell_text, bold=(i == 0))
    doc.add_paragraph()


def write_docx(pages: list[list[Block]], path: Path, page_breaks: bool = True) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    # Fallback font for CJK text (e.g. untranslated names or zh target)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")

    for n, blocks in enumerate(pages):
        if n and page_breaks:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        for b in blocks:
            if b.kind == "heading":
                _add_inline(doc.add_heading(level=min(b.level, 4)), b.text)
            elif b.kind == "para":
                _add_inline(doc.add_paragraph(), b.text)
            elif b.kind == "list":
                numbered = b.prefix.strip()[:1].isdigit()
                style = "List Number" if numbered else "List Bullet"
                if b.level >= 1:
                    style += " 2"
                _add_inline(doc.add_paragraph(style=style), b.text)
            elif b.kind == "table":
                _add_table(doc, b.rows)
            elif b.kind == "html_table":
                _add_table(doc, html_table_rows(b.raw))
            elif b.kind == "code":
                lines = b.raw.split("\n")[1:]  # drop opening fence with language tag
                if lines and lines[-1].strip().startswith(("```", "~~~")):
                    lines = lines[:-1]
                run = doc.add_paragraph().add_run("\n".join(lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif b.kind == "math":
                run = doc.add_paragraph().add_run(b.raw.strip().strip("$").removeprefix("\\[").removesuffix("\\]").strip())
                run.font.name = "Cambria Math"
            # raw (hr, image placeholders, comments) is not rendered
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
